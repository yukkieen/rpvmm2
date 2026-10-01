from __future__ import annotations

import csv
import html
import json
import re
import unicodedata
import warnings
from pathlib import Path
from urllib.parse import parse_qs, unquote

from bs4 import BeautifulSoup, Tag
from openpyxl import load_workbook

DATA_DIR = Path(__file__).resolve().parent.parent / "data"
DATA_PATH = DATA_DIR / "professors.json"
LATTES_DIR = DATA_DIR / "lattes"
ARTICLES_PATH = DATA_DIR / "artigos.csv"
SCOPUS_PATH = DATA_DIR / "scopus_fwci.csv"
QUALIS_PATH = Path(__file__).resolve().parent.parent / "qualis_21-24.xlsx"
ARTICLE_FIELDS = ["professor", "artigo", "doi", "citations", "fwci", "source", "qualis"]

YEAR_MIN = 2025
YEAR_MAX = 2028

DOI_HREF_RE = re.compile(
    r"(?:https?://(?:dx\.)?doi\.org/)?(10\.\d{4,9}/[^\s\"'<>]+)",
    re.IGNORECASE,
)
DOI_TEXT_RE = re.compile(r"\b(10\.\d{4,9}/[-._;()/:A-Z0-9]+)", re.IGNORECASE)
PAREN_RE = re.compile(r"\([^)]*\)")
NON_ALNUM_RE = re.compile(r"[^a-z0-9]+")
ESTRATO_RANK = {
    "A1": 0,
    "A2": 1,
    "A3": 2,
    "A4": 3,
    "B1": 4,
    "B2": 5,
    "B3": 6,
    "B4": 7,
    "B5": 8,
    "C": 9,
}


def load_professors(path: Path = DATA_PATH) -> list[dict]:
    return json.loads(path.read_text(encoding="utf-8"))


def unique_professors(professors: list[dict]) -> list[dict]:
    """Duplicate listings on the site keep the first occurrence of each name."""
    by_name: dict[str, dict] = {}
    for professor in professors:
        name = professor.get("name")
        if not name:
            continue
        by_name.setdefault(name.casefold(), professor)
    return list(by_name.values())


def cv_name(html_path: Path) -> str:
    soup = BeautifulSoup(html_path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    el = soup.select_one("h2.nome")
    return el.get_text(strip=True) if el else ""


def lattes_path_for(professor: dict) -> Path | None:
    if professor.get("lattes_html"):
        path = Path(professor["lattes_html"])
        if path.exists():
            return path
    lid = professor.get("lattes_id")
    if lid:
        path = LATTES_DIR / f"{lid}.html"
        if path.exists():
            return path

    name_key = (professor.get("name") or "").casefold()
    if not name_key or not LATTES_DIR.exists():
        return None
    for path in LATTES_DIR.glob("*.html"):
        if cv_name(path).casefold() == name_key:
            return path
    return None


def clean_text(value: str | None) -> str:
    if not value:
        return ""
    text = html.unescape(unquote(value))
    return re.sub(r"\s+", " ", text).strip()


def normalize_title(title: str) -> str:
    text = clean_text(title).casefold()
    text = text.replace("´", "'").replace("’", "'").replace("`", "'")
    return text


def normalize_doi(value: str | None) -> str:
    if not value:
        return ""
    text = html.unescape(unquote(value)).strip()
    if not text:
        return ""
    match = DOI_HREF_RE.search(text)
    doi = match.group(1) if match else text
    doi = doi.rstrip(").,;")
    return doi


def normalize_source(name: str | None) -> str:
    if not name:
        return ""
    text = unicodedata.normalize("NFKD", name)
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = PAREN_RE.sub(" ", text)
    text = text.casefold().replace("&", " and ")
    text = NON_ALNUM_RE.sub(" ", text)
    return " ".join(text.split())


def _better_estrato(current: str, candidate: str) -> bool:
    return ESTRATO_RANK.get(candidate, 99) < ESTRATO_RANK.get(current, 99)


def load_qualis(path: Path = QUALIS_PATH) -> dict[str, str]:
    """Map normalized journal title → best CAPES Qualis estrato (2021–2024)."""
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Workbook contains no default style")
        wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    by_title: dict[str, str] = {}
    seen_issn: set[str] = set()
    # Sucupira xlsx reports wrong sheet dimensions; pin the used range.
    rows = ws.iter_rows(min_row=2, max_row=500_000, max_col=4, values_only=True)
    for row in rows:
        if not row or not any(row):
            continue
        issn = str(row[0] or "").strip()
        title = str(row[1] or "").strip()
        estrato = str(row[3] or "").strip().upper()
        if issn:
            if issn in seen_issn:
                continue
            seen_issn.add(issn)
        if not title or not estrato:
            continue
        key = normalize_source(title)
        if not key:
            continue
        if key not in by_title or _better_estrato(by_title[key], estrato):
            by_title[key] = estrato
    wb.close()
    return by_title


def qualis_of(source: str, by_title: dict[str, str]) -> str:
    return by_title.get(normalize_source(source), "")


def year_of(article: Tag) -> int | None:
    el = article.select_one('span.informacao-artigo[data-tipo-ordenacao="ano"]')
    if not el:
        return None
    raw = el.get_text(strip=True)
    if raw.isdigit():
        return int(raw)
    return None


def title_from_cvuri(article: Tag) -> str:
    citado = article.select_one("span.citado[cvuri]")
    if not citado:
        return ""
    qs = parse_qs(citado["cvuri"].split("?", 1)[-1])
    return clean_text((qs.get("titulo") or [""])[0])


def title_from_text(article: Tag) -> str:
    """Lattes citation: AUTHORS . Title. JOURNAL, ..."""
    transform = article.select_one("span.transform")
    if not transform:
        return ""
    clone = BeautifulSoup(str(transform), "html.parser")
    root = clone.select_one("span.transform") or clone
    for junk in root.select("span.informacao-artigo, a.icone-producao, span.citado, sup"):
        junk.decompose()
    text = clean_text(root.get_text(" ", strip=True))
    match = re.search(r"\s\.\s(.+?)\.\s", text)
    return match.group(1).strip() if match else ""


def doi_of(article: Tag) -> str:
    link = article.select_one("a.icone-doi[href]")
    if link:
        doi = normalize_doi(link.get("href"))
        if doi:
            return doi

    citado = article.select_one("span.citado[cvuri]")
    if citado:
        qs = parse_qs(citado["cvuri"].split("?", 1)[-1])
        doi = normalize_doi((qs.get("doi") or [""])[0])
        if doi:
            return doi

    text = article.get_text(" ", strip=True)
    match = DOI_TEXT_RE.search(text)
    return normalize_doi(match.group(1) if match else "")


def parse_article(article: Tag) -> dict | None:
    year = year_of(article)
    if year is None or year < YEAR_MIN or year > YEAR_MAX:
        return None

    title = title_from_cvuri(article) or title_from_text(article)
    if not title:
        return None

    return {"artigo": title, "doi": doi_of(article), "ano": year}


def parse_cv(html_path: Path) -> list[dict]:
    soup = BeautifulSoup(html_path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    seen_titles: set[str] = set()
    seen_dois: set[str] = set()
    rows: list[dict] = []
    for article in soup.select("div.artigo-completo"):
        parsed = parse_article(article)
        if not parsed:
            continue
        title_key = normalize_title(parsed["artigo"])
        doi_key = parsed["doi"].casefold()
        if title_key in seen_titles or (doi_key and doi_key in seen_dois):
            continue
        seen_titles.add(title_key)
        if doi_key:
            seen_dois.add(doi_key)
        rows.append(parsed)
    return rows


def fill_missing_dois(rows: list[dict]) -> None:
    """Reuse a DOI when a coauthor in the same set already listed it."""
    by_title: dict[str, str] = {}
    for row in rows:
        if row["doi"]:
            by_title.setdefault(normalize_title(row["artigo"]), row["doi"])
    for row in rows:
        if not row["doi"]:
            row["doi"] = by_title.get(normalize_title(row["artigo"]), "")


def collect_articles(professors: list[dict]) -> list[dict]:
    rows: list[dict] = []
    for professor in unique_professors(professors):
        path = lattes_path_for(professor)
        if path is None:
            print(f"aviso: sem HTML do Lattes para {professor['name']}")
            continue

        for article in parse_cv(path):
            rows.append(
                {
                    "professor": professor["name"],
                    "artigo": article["artigo"],
                    "doi": article["doi"],
                }
            )
    fill_missing_dois(rows)
    rows.sort(key=lambda r: (r["professor"].casefold(), r["artigo"].casefold()))
    return rows


def load_articles(path: Path = ARTICLES_PATH) -> list[dict]:
    with path.open(encoding="utf-8", newline="") as fh:
        return list(csv.DictReader(fh))


def load_scopus_metrics(path: Path = SCOPUS_PATH) -> dict[str, dict]:
    by_doi: dict[str, dict] = {}
    with path.open(encoding="utf-8", newline="") as fh:
        for row in csv.DictReader(fh):
            doi = normalize_doi(row.get("doi"))
            if not doi:
                continue
            by_doi[doi.casefold()] = {
                "citations": (row.get("citations") or "").strip(),
                "fwci": (row.get("fwci") or "").strip(),
                "source": (row.get("source") or "").strip(),
            }
    return by_doi


def merge_scopus_metrics(
    rows: list[dict],
    metrics: dict[str, dict] | None = None,
    scopus_path: Path = SCOPUS_PATH,
    qualis_path: Path = QUALIS_PATH,
) -> list[dict]:
    """Attach Scopus citations, FWCI, source, and CAPES Qualis to each article by DOI."""
    if metrics is None:
        metrics = load_scopus_metrics(scopus_path) if scopus_path.exists() else {}
    by_title = load_qualis(qualis_path) if qualis_path.exists() else {}
    for row in rows:
        hit = metrics.get(normalize_doi(row.get("doi")).casefold(), {})
        row["citations"] = hit.get("citations", "")
        row["fwci"] = hit.get("fwci", "")
        source = hit.get("source", "")
        row["source"] = source
        row["qualis"] = qualis_of(source, by_title) if source else ""
    return rows


def save_articles(rows: list[dict], path: Path = ARTICLES_PATH) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=ARTICLE_FIELDS, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return path


def merge_scopus_into_articles(
    articles_path: Path = ARTICLES_PATH,
    scopus_path: Path = SCOPUS_PATH,
) -> Path:
    rows = merge_scopus_metrics(load_articles(articles_path), scopus_path=scopus_path)
    return save_articles(rows, articles_path)


if __name__ == "__main__":
    professors = load_professors()
    rows = collect_articles(professors)
    path = save_articles(rows)
    with_doi = sum(1 for r in rows if r["doi"])
    print(f"Wrote {len(rows)} artigos ({with_doi} com DOI) to {path}")
    path = merge_scopus_into_articles()
    merged = load_articles(path)
    with_fwci = sum(1 for r in merged if r.get("fwci"))
    with_qualis = sum(1 for r in merged if r.get("qualis"))
    print(f"Merged Scopus metrics into {path} ({with_fwci} com FWCI, {with_qualis} com Qualis)")
